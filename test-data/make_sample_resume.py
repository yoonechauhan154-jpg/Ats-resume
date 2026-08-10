"""Generate a realistic sample resume (.docx and .pdf) for integration testing.

Fields intentionally included so every parser path is exercised:
contact block, summary, experience with metric-heavy bullets, education, skills.
"""

from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches

OUT_DIR = Path("/workspace/test-data")
OUT_DIR.mkdir(exist_ok=True)

RESUME_LINES = [
    "Alex Morgan",
    "Senior Backend Engineer",
    "alex.morgan@example.com | +1 (555) 010-2030 | Seattle, WA | linkedin.com/in/alexmorgan",
    "",
    "SUMMARY",
    "Backend engineer with 7 years of experience designing distributed systems and data pipelines. "
    "Strong track record of improving reliability and performance for high-traffic services.",
    "",
    "EXPERIENCE",
    "Senior Backend Engineer, Northwind Labs (2020 - Present)",
    "- Led the rebuild of the payments API, cutting p95 latency by 38%",
    "- Built event-driven services in Python that process 2M messages/day",
    "- Reduced deploy time by 40% by automating tests and CI/CD",
    "- Mentored 5 junior engineers and introduced code review standards",
    "",
    "Backend Engineer, Cloudline Systems (2017 - 2020)",
    "- Designed REST APIs for a B2B analytics platform serving 500k users",
    "- Improved database query performance with indexing, cutting report time by 55%",
    "- Migrated 12 microservices from on-premise to AWS",
    "- On-call incident response reduced by 30% through monitoring improvements",
    "",
    "Software Engineer, Startup Co (2015 - 2017)",
    "- Shipped features across a Python/PostgreSQL stack",
    "- Wrote unit tests raising coverage from 41% to 78%",
    "",
    "EDUCATION",
    "B.Sc. Computer Science, University of Washington (2015)",
    "",
    "SKILLS",
    "Python, SQL, PostgreSQL, AWS, Docker, Git, Linux, REST APIs, unit testing",
    "",
    "CERTIFICATIONS",
    "AWS Certified Solutions Architect - Associate (2022)",
]


def build_docx(path: Path) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for line in RESUME_LINES:
        stripped = line.strip()
        if not stripped:
            doc.add_paragraph("")
        elif stripped.isupper():
            p = doc.add_paragraph()
            run = p.add_run(stripped)
            run.bold = True
        elif stripped.startswith("-"):
            p = doc.add_paragraph(stripped[1:].strip(), style="List Bullet")
        else:
            p = doc.add_paragraph()
            if any(role in stripped for role in ("Northwind", "Cloudline", "Startup")):
                p.add_run(stripped).bold = True
            else:
                p.add_run(stripped)
    doc.save(str(path))


def build_pdf(path: Path) -> None:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    y = 60
    for line in RESUME_LINES:
        if page.rect.height - y < 50:
            page = doc.new_page()
            y = 60
        page.insert_text((60, y), line, fontsize=10.5, fontname="helv")
        y += 16
    doc.save(str(path))
    doc.close()


if __name__ == "__main__":
    build_docx(OUT_DIR / "sample_resume.docx")
    build_pdf(OUT_DIR / "sample_resume.pdf")
    print("Wrote", OUT_DIR / "sample_resume.docx", "and", OUT_DIR / "sample_resume.pdf")
