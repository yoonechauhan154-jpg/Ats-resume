"""Tests for format checking and parsing."""

from pathlib import Path

import pytest

from app.services.parser.base import ParsedText
from app.services.scoring.format_checker import analyze_format
from app.services.parser.pdf_parser import parse_pdf
from app.services.parser.docx_parser import parse_docx


def test_format_scanned_critical():
    issues, score = analyze_format(
        ParsedText(text="", scanned=True, has_text_layer=False), 0
    )
    assert score < 60
    assert any(i.type.value == "scanned" and i.severity.value == "critical" for i in issues)


def test_format_tables_detected():
    issues, score = analyze_format(
        ParsedText(text="x", has_tables=True, table_count=3, has_text_layer=True), 300
    )
    assert score < 100
    assert any(i.type.value == "table" for i in issues)


def test_format_missing_contact():
    issues, _ = analyze_format(
        ParsedText(text="John Smith", has_text_layer=True), 300
    )
    assert any(i.type.value == "missing_contact" for i in issues)


def test_format_clean_resume_high_score():
    issues, score = analyze_format(
        ParsedText(
            text="John Smith\njohn@example.com\n+1 555 123 4567\nPython developer.",
            has_text_layer=True,
            fonts_used=["arial"],
            font_size_min=10,
        ),
        320,
    )
    assert score >= 95
    assert not issues


def _make_test_pdf(tmp_path: Path) -> Path:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "John Smith", fontsize=14, fontname="helv")
    page.insert_text((72, 96), "Senior Software Engineer", fontsize=11, fontname="helv")
    page.insert_text((72, 130), "Built REST APIs with Python and PostgreSQL.", fontsize=11, fontname="helv")
    path = tmp_path / "resume.pdf"
    doc.save(str(path))
    doc.close()
    return path


def test_pdf_parser_extracts_text(tmp_path):
    path = _make_test_pdf(tmp_path)
    parsed = parse_pdf(path)
    assert parsed.has_text_layer
    assert "Python" in parsed.text
    assert parsed.word_count > 5


def test_docx_parser(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_heading("John Smith", 0)
    doc.add_paragraph("Senior Software Engineer")
    doc.add_paragraph("Built REST APIs with Python and PostgreSQL.")
    doc.add_paragraph("Skills: Python, React, AWS")
    path = tmp_path / "resume.docx"
    doc.save(str(path))

    parsed = parse_docx(path)
    assert "Python" in parsed.text
    assert "PostgreSQL" in parsed.text
    assert parsed.has_text_layer


def _tesseract_ok() -> bool:
    try:
        import pytesseract  # noqa: F401
        from PIL import Image  # noqa: F401
        return True
    except Exception:
        return False


@pytest.mark.skipif(
    not _tesseract_ok(),
    reason="tesseract not installed",
)
def test_pdf_parser_ocr_scanned(tmp_path):
    import fitz
    from PIL import Image
    import io

    img = Image.new("RGB", (600, 200), "white")
    from PIL import ImageDraw

    d = ImageDraw.Draw(img)
    d.text((20, 50), "John Smith - Python Engineer", fill="black")
    buf = io.BytesIO()
    img.save(buf, format="PNG")

    doc = fitz.open()
    page = doc.new_page()
    page.insert_image(page.rect, stream=buf.getvalue())
    path = tmp_path / "scanned.pdf"
    doc.save(str(path))
    doc.close()

    parsed = parse_pdf(path)
    assert parsed.scanned or parsed.was_ocr
